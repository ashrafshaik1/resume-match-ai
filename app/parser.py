from pypdf import PdfReader
from docx import Document


# --------------------------------
# Extract text from PDF
# --------------------------------

def extract_text_from_pdf(file_path):

    reader = PdfReader(file_path)

    text = []

    for page in reader.pages:

        page_text = page.extract_text()

        if page_text:
            text.append(page_text)

    return "\n".join(text)


# --------------------------------
# Extract text from DOCX
# --------------------------------

def extract_text_from_docx(file_path):

    document = Document(file_path)

    text = []

    # Extract paragraphs
    for paragraph in document.paragraphs:

        if paragraph.text.strip():

            text.append(
                paragraph.text.strip()
            )

    # Extract text from tables
    for table in document.tables:

        for row in table.rows:

            for cell in row.cells:

                if cell.text.strip():

                    text.append(
                        cell.text.strip()
                    )

    return "\n".join(text)